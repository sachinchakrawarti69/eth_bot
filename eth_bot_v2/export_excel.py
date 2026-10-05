"""Export OHLCV to formatted Excel."""

from __future__ import annotations
import sqlite3
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, StockChart
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT    = Path(__file__).resolve().parent
DB      = ROOT / "database" / "market.db"
OUT_DIR = ROOT / "exports"
OUT     = OUT_DIR / "ETHUSDT_1y_ohlcv.xlsx"

SYMBOL, TF, YEARS = "ETHUSDT", "1d", 1

COLOR_HEADER_BG = "1F4E78"
COLOR_HEADER_FG = "FFFFFF"
COLOR_UP        = "C6EFCE"
COLOR_DOWN      = "FFC7CE"
COLOR_UP_TEXT   = "006100"
COLOR_DOWN_TEXT = "9C0006"
COLOR_BAND      = "F2F2F2"
_thin  = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def load():
    with sqlite3.connect(DB) as c:
        df = pd.read_sql_query(
            "SELECT open_time, open, high, low, close, volume "
            "FROM ohlcv WHERE symbol=? AND timeframe=? ORDER BY open_time",
            c, params=(SYMBOL, TF))
    df["date"] = pd.to_datetime(df["open_time"], unit="ms", utc=True).dt.tz_convert(None)
    cutoff = df["date"].max() - pd.Timedelta(days=365*YEARS)
    df = df[df["date"] >= cutoff].copy()
    df["return"]    = df["close"].pct_change()
    df["range"]     = df["high"] - df["low"]
    df["sma_20"]    = df["close"].rolling(20).mean()
    df["sma_50"]    = df["close"].rolling(50).mean()
    df["ema_12"]    = df["close"].ewm(span=12, adjust=False).mean()
    df["ema_26"]    = df["close"].ewm(span=26, adjust=False).mean()
    df["macd"]      = df["ema_12"] - df["ema_26"]
    df["macd_sig"]  = df["macd"].ewm(span=9, adjust=False).mean()
    df["macd_hist"] = df["macd"] - df["macd_sig"]
    d  = df["close"].diff()
    up = d.clip(lower=0).rolling(14).mean()
    dn = (-d.clip(upper=0)).rolling(14).mean()
    df["rsi_14"] = 100 - 100 / (1 + up / dn.replace(0, np.nan))
    df["dir"] = np.where(df["close"] >= df["open"], "up", "down")
    return df.drop(columns=["open_time"]).set_index("date")


def style_header(ws, ncols=None):
    ncols = ncols or ws.max_column
    fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    font = Font(color=COLOR_HEADER_FG, bold=True)
    for c in range(1, ncols+1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill; cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=2, column=1)


def autosize(ws, min_w=10, max_w=22):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, longest+2))


def sheet_candles(wb, df):
    ws = wb.create_sheet("Candles")
    ws.append(["Date","Open","High","Low","Close","Volume","Return","Range","Dir"])
    for d, row in df.iterrows():
        ws.append([d.date(),
                   float(row.open), float(row.high), float(row.low), float(row.close),
                   float(row.volume),
                   float(row["return"]) if pd.notna(row["return"]) else None,
                   float(row["range"]), row["dir"]])
    style_header(ws)
    n = ws.max_row
    for r in range(2, n+1):
        for c in (2,3,4,5): ws.cell(r,c).number_format = "#,##0.00"
        ws.cell(r,6).number_format = "#,##0"
        ws.cell(r,7).number_format = "+0.00%;-0.00%"
        ws.cell(r,8).number_format = "#,##0.00"
    ws.conditional_formatting.add(f"A2:I{n}",
        FormulaRule(formula=['$I2="up"'], fill=PatternFill("solid", fgColor=COLOR_UP)))
    ws.conditional_formatting.add(f"A2:I{n}",
        FormulaRule(formula=['$I2="down"'], fill=PatternFill("solid", fgColor=COLOR_DOWN)))
    ws.conditional_formatting.add(f"G2:G{n}",
        ColorScaleRule(start_type="min", start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="max", end_color="63BE7B"))
    autosize(ws)
    ws.auto_filter.ref = f"A1:I{n}"
    ws.sheet_view.showGridLines = False


def sheet_indicators(wb, df):
    ws = wb.create_sheet("Indicators")
    ws.append(["Date","Close","SMA20","SMA50","EMA12","EMA26",
               "MACD","MACD_Sig","MACD_Hist","RSI14"])
    for d, row in df.iterrows():
        ws.append([d.date(),
                   float(row.close) if pd.notna(row.close) else None,
                   float(row.sma_20) if pd.notna(row.sma_20) else None,
                   float(row.sma_50) if pd.notna(row.sma_50) else None,
                   float(row.ema_12) if pd.notna(row.ema_12) else None,
                   float(row.ema_26) if pd.notna(row.ema_26) else None,
                   float(row["macd"]) if pd.notna(row["macd"]) else None,
                   float(row["macd_sig"]) if pd.notna(row["macd_sig"]) else None,
                   float(row["macd_hist"]) if pd.notna(row["macd_hist"]) else None,
                   float(row["rsi_14"]) if pd.notna(row["rsi_14"]) else None])
    style_header(ws)
    n = ws.max_row
    for r in range(2, n+1):
        for c in range(2, 11): ws.cell(r,c).number_format = "#,##0.00"
    ws.conditional_formatting.add(f"J2:J{n}",
        CellIsRule(operator="greaterThan", formula=["70"],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"J2:J{n}",
        CellIsRule(operator="lessThan", formula=["30"],
                   fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"I2:I{n}",
        ColorScaleRule(start_type="min", start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="max", end_color="63BE7B"))
    autosize(ws)
    chart = LineChart()
    chart.title = f"{SYMBOL} close / SMA20 / SMA50"
    chart.height = 8; chart.width = 22
    data = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "L2")


def sheet_monthly(wb, df):
    ws = wb.create_sheet("Monthly")
    r = df["return"].dropna()
    monthly = (1 + r).groupby([r.index.year, r.index.month]).prod() - 1
    monthly.index.names = ["year","month"]
    pivot = monthly.unstack("month")
    labels = ["Jan","Feb","Mar","Apr","May","Jun",
              "Jul","Aug","Sep","Oct","Nov","Dec"]
    pivot.columns = [labels[c-1] for c in pivot.columns]
    ws.append(["Year"] + list(pivot.columns))
    for year, row in pivot.iterrows():
        ws.append([int(year)] + [float(v) if pd.notna(v) else None for v in row])
    style_header(ws)
    n = ws.max_row; ncols = len(pivot.columns)
    for r_ in range(2, n+1):
        for c in range(2, 2+ncols): ws.cell(r_,c).number_format = "+0.0%;-0.0%;0.0%"
    rng = f"B2:{get_column_letter(1+ncols)}{n}"
    ws.conditional_formatting.add(rng,
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor="C6EFCE"),
                   font=Font(color=COLOR_UP_TEXT, bold=True)))
    ws.conditional_formatting.add(rng,
        CellIsRule(operator="lessThan", formula=["0"],
                   fill=PatternFill("solid", fgColor="FFC7CE"),
                   font=Font(color=COLOR_DOWN_TEXT, bold=True)))
    autosize(ws)
    ws.sheet_view.showGridLines = False


def sheet_stats(wb, df):
    ws = wb.create_sheet("Stats")
    r   = df["return"].dropna()
    cum = (1 + r).cumprod()
    dd  = (cum / cum.cummax() - 1).min()
    rows = [
        ("Symbol", SYMBOL),
        ("Period", str(df.index.min().date()) + " to " + str(df.index.max().date())),
        ("Days", len(df)),
        ("Start price", float(df.close.iloc[0])),
        ("End price", float(df.close.iloc[-1])),
        ("Total return", float(cum.iloc[-1] - 1)),
        ("Annualized vol", float(r.std(ddof=0) * np.sqrt(365))),
        ("Sharpe (0 rf)", float(r.mean() / r.std(ddof=0) * np.sqrt(365))),
        ("Max drawdown", float(dd)),
        ("Best day", float(r.max())),
        ("Worst day", float(r.min())),
        ("Avg daily volume", float(df.volume.mean())),
        ("Max close", float(df.close.max())),
        ("Min close", float(df.close.min())),
    ]
    ws.append(["Metric","Value"])
    for k,v in rows: ws.append([k,v])
    style_header(ws)
    for r_ in range(2, ws.max_row+1):
        k = ws.cell(r_,1).value; v = ws.cell(r_,2).value; c = ws.cell(r_,2)
        if k in ("Start price","End price","Max close","Min close"): c.number_format="#,##0.00"
        elif k in ("Total return","Annualized vol","Max drawdown","Best day","Worst day"): c.number_format="+0.00%;-0.00%"
        elif k == "Sharpe (0 rf)": c.number_format = "0.00"
        elif k == "Avg daily volume": c.number_format = "#,##0"
        if k in ("Total return","Best day") and isinstance(v,(int,float)):
            c.fill = PatternFill("solid", fgColor=COLOR_UP)
        elif k in ("Max drawdown","Worst day") and isinstance(v,(int,float)):
            c.fill = PatternFill("solid", fgColor=COLOR_DOWN)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 26
    ws.sheet_view.showGridLines = False


def sheet_candlestick(wb, df):
    ws = wb.create_sheet("Candlestick")
    ws.append(["Date","Open","High","Low","Close"])
    for d, row in df.iterrows():
        ws.append([d.date(), float(row.open), float(row.high),
                   float(row.low), float(row.close)])
    style_header(ws)
    n = ws.max_row
    autosize(ws)
    chart = StockChart()
    chart.title = f"{SYMBOL} - last {YEARS}y"
    chart.height = 14; chart.width = 30
    data = Reference(ws, min_col=2, max_col=5, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "G2")


def main():
    if not DB.exists():
        print("DB not found:", DB)
        return 1
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    df = load()
    print(f"loaded {len(df)} rows  {df.index.min().date()} -> {df.index.max().date()}")
    wb = Workbook()
    wb.remove(wb.active)
    sheet_candles(wb, df)
    sheet_indicators(wb, df)
    sheet_monthly(wb, df)
    sheet_stats(wb, df)
    sheet_candlestick(wb, df)
    wb.save(OUT)
    print("wrote", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
