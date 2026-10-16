"""
excel_update.py - build a single formatted Excel workbook from every data source.

Location: D:\\Projects\\done\\eth_bot\\eth_bot_v2\\excel_update.py

Usage:
    python excel_update.py                       # full profile
    python excel_update.py --profile ohlcv
    python excel_update.py --profile forecasts
    python excel_update.py --profile ml
    python excel_update.py --profile quality
    python excel_update.py --profile options
    python excel_update.py --sheet stats --sheet monthly
    python excel_update.py --split               # one file per sheet
    python excel_update.py --list
    python excel_update.py --out custom_name.xlsx
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
import time
from datetime import date, datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

try:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("pip install openpyxl")
    sys.exit(1)


# ============================================================
# paths
# ============================================================

ROOT        = Path(__file__).resolve().parent
DB_MARKET   = ROOT / "database" / "market.db"
DB_OPTIONS  = ROOT / "database" / "options.db"
FORECASTS   = ROOT / "reports" / "forecasts"
PREDICTIONS = ROOT / "reports" / "predictions"
BACKTESTS   = ROOT / "reports" / "backtest6y"
ML_RUNS     = ROOT / "ml" / "_runs"
OUT_DIR     = ROOT / "exports"

TODAY  = date.today().isoformat()
SYMBOL = "ETHUSDT"
TF     = "1d"


# ============================================================
# styling
# ============================================================

HDR_BG    = "1F4E78"
HDR_FG    = "FFFFFF"
UP_BG     = "C6EFCE"
DOWN_BG   = "FFC7CE"
UP_TEXT   = "006100"
DOWN_TEXT = "9C0006"
BAND_BG   = "F2F2F2"

_thin  = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

FMT_PRICE = "#,##0.00"
FMT_PCT   = "+0.00%;-0.00%"
FMT_INT   = "#,##0"
FMT_IV    = "0.00%"


# ============================================================
# db helpers
# ============================================================

def q_market(sql: str, **p) -> pd.DataFrame:
    if not DB_MARKET.exists():
        return pd.DataFrame()
    with sqlite3.connect(DB_MARKET) as c:
        return pd.read_sql_query(sql, c, params=p)


def q_options(sql: str, **p) -> pd.DataFrame:
    if not DB_OPTIONS.exists():
        return pd.DataFrame()
    with sqlite3.connect(DB_OPTIONS) as c:
        return pd.read_sql_query(sql, c, params=p)


# ============================================================
# sources
# ============================================================

def src_ohlcv() -> pd.DataFrame:
    df = q_market(
        "SELECT open_time, open, high, low, close, volume "
        "FROM ohlcv WHERE symbol=:s AND timeframe=:t ORDER BY open_time",
        s=SYMBOL, t=TF)
    if df.empty:
        return df
    df["date"] = pd.to_datetime(df["open_time"], unit="ms", utc=True).dt.tz_convert(None)
    df["return"] = df["close"].pct_change()
    df["range"]  = df["high"] - df["low"]
    df["dir"]    = np.where(df["close"] >= df["open"], "up", "down")
    return df.drop(columns=["open_time"]).set_index("date")


def src_indicators() -> pd.DataFrame:
    df = src_ohlcv()
    if df.empty:
        return df
    df = df.copy()
    df["sma_20"]    = df["close"].rolling(20).mean()
    df["sma_50"]    = df["close"].rolling(50).mean()
    df["sma_200"]   = df["close"].rolling(200).mean()
    df["ema_12"]    = df["close"].ewm(span=12, adjust=False).mean()
    df["ema_26"]    = df["close"].ewm(span=26, adjust=False).mean()
    df["macd"]      = df["ema_12"] - df["ema_26"]
    df["macd_sig"]  = df["macd"].ewm(span=9, adjust=False).mean()
    df["macd_hist"] = df["macd"] - df["macd_sig"]
    d  = df["close"].diff()
    up = d.clip(lower=0).rolling(14).mean()
    dn = (-d.clip(upper=0)).rolling(14).mean()
    df["rsi_14"] = 100 - 100 / (1 + up / dn.replace(0, np.nan))
    tr = pd.concat([
        df["high"] - df["low"],
        (df["high"] - df["close"].shift()).abs(),
        (df["low"]  - df["close"].shift()).abs(),
    ], axis=1).max(axis=1)
    df["atr_14"] = tr.rolling(14).mean()
    return df


def src_monthly() -> pd.DataFrame:
    df = src_ohlcv()
    if df.empty:
        return df
    r = df["return"].dropna()
    monthly = (1 + r).groupby([r.index.year, r.index.month]).prod() - 1
    monthly.index.names = ["year", "month"]
    pivot = monthly.unstack("month")
    labels = ["Jan","Feb","Mar","Apr","May","Jun",
              "Jul","Aug","Sep","Oct","Nov","Dec"]
    pivot.columns = [labels[int(c)-1] for c in pivot.columns]
    return pivot.reset_index()


def src_stats() -> pd.DataFrame:
    df = src_ohlcv()
    if df.empty:
        return df
    r = df["return"].dropna()
    cum = (1 + r).cumprod()
    dd = (cum / cum.cummax() - 1).min()
    rows = [
        ("Symbol", SYMBOL),
        ("Period", f"{df.index.min().date()} to {df.index.max().date()}"),
        ("Days", len(df)),
        ("Start price", float(df.close.iloc[0])),
        ("End price", float(df.close.iloc[-1])),
        ("Total return", float(cum.iloc[-1] - 1)),
        ("Annualized vol", float(r.std(ddof=0) * np.sqrt(365))),
        ("Sharpe (0 rf)", float(r.mean() / r.std(ddof=0) * np.sqrt(365))),
        ("Max drawdown", float(dd)),
        ("Best day", float(r.max())),
        ("Worst day", float(r.min())),
        ("Avg daily volume", float(df.volume.mean())),
        ("Max close", float(df.close.max())),
        ("Min close", float(df.close.min())),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def src_funding() -> pd.DataFrame:
    df = q_market("SELECT ts, rate FROM funding_rate ORDER BY ts")
    if df.empty: return df
    df["date"] = pd.to_datetime(df["ts"], unit="ms", utc=True).dt.tz_convert(None).dt.date
    return df[["date", "rate"]]


def src_open_interest() -> pd.DataFrame:
    df = q_market("SELECT ts, oi, oi_value FROM open_interest ORDER BY ts")
    if df.empty: return df
    df["date"] = pd.to_datetime(df["ts"], unit="ms", utc=True).dt.tz_convert(None).dt.date
    return df[["date", "oi", "oi_value"]]


def src_fear_greed() -> pd.DataFrame:
    df = q_market("SELECT ts, value, label FROM fear_greed ORDER BY ts")
    if df.empty: return df
    df["date"] = pd.to_datetime(df["ts"], unit="ms", utc=True).dt.tz_convert(None).dt.date
    return df[["date", "value", "label"]]


def src_option_chain() -> pd.DataFrame:
    df = q_options("SELECT * FROM option_chain ORDER BY expiry, strike LIMIT 2000")
    if df.empty:
        df = q_market("SELECT * FROM option_chain ORDER BY expiry, strike LIMIT 2000")
    if df.empty: return df
    if "expiry" in df.columns:
        df["expiry"] = pd.to_datetime(df["expiry"], unit="ms", utc=True,
                                      errors="coerce").dt.tz_convert(None).dt.date
    if "ts" in df.columns:
        df["ts"] = pd.to_datetime(df["ts"], unit="ms", utc=True,
                                  errors="coerce").dt.tz_convert(None).dt.date
    return df


def src_inventory() -> pd.DataFrame:
    rows = []
    for path, label in ((DB_MARKET, "market"), (DB_OPTIONS, "options")):
        if not path.exists():
            continue
        with sqlite3.connect(path) as c:
            tables = [r[0] for r in c.execute(
                "SELECT name FROM sqlite_master WHERE type='table' "
                "AND name NOT LIKE 'sqlite_%' ORDER BY name")]
            for t in tables:
                n = c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                rows.append((label, t, int(n)))
    return pd.DataFrame(rows, columns=["db", "table", "rows"])


def src_forecast_dashboard() -> pd.DataFrame:
    if not FORECASTS.exists():
        return pd.DataFrame()
    rows = []
    for f in sorted(FORECASTS.glob("*.json")):
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue

        entries = []
        if isinstance(d, list):
            entries = d
        elif isinstance(d, dict):
            # a dict with a nested list?
            for k in ("forecasts", "signals", "rows", "items"):
                if k in d and isinstance(d[k], list):
                    entries = d[k]
                    break
            else:
                entries = [d]

        for item in entries:
            if not isinstance(item, dict):
                continue
            rows.append({
                "file":       f.name,
                "symbol":     item.get("symbol", SYMBOL),
                "action":     item.get("action", item.get("direction", "")),
                "price":      item.get("price", item.get("last_close", "")),
                "target":     item.get("pred_close", item.get("target", "")),
                "pred_ret":   item.get("pred_ret", ""),
                "confidence": item.get("confidence",
                                       item.get("dir_acc",
                                                item.get("dir_accuracy", ""))),
                "generated":  item.get("as_of", item.get("generated_at", "")),
            })
    return pd.DataFrame(rows)


def src_ml_leaderboard() -> pd.DataFrame:
    if not ML_RUNS.exists():
        return pd.DataFrame()
    rows = []
    for m in ML_RUNS.glob("*/metrics.json"):
        try:
            d = json.loads(m.read_text(encoding="utf-8"))
        except Exception:
            continue
        if not isinstance(d, dict):
            continue
        rows.append({
            "run":       m.parent.name,
            "model":     d.get("model", ""),
            "symbol":    d.get("symbol", ""),
            "accuracy":  d.get("accuracy", d.get("acc", np.nan)),
            "precision": d.get("precision", np.nan),
            "recall":    d.get("recall", np.nan),
            "f1":        d.get("f1", np.nan),
            "n_test":    d.get("n_test", d.get("n", np.nan)),
        })
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values("accuracy", ascending=False, na_position="last")
    return df


def src_backtest() -> pd.DataFrame:
    """Read the newest reports/backtest6y/*.json and flatten it."""
    if not BACKTESTS.exists():
        return pd.DataFrame()
    files = sorted(BACKTESTS.glob("*.json"))
    if not files:
        return pd.DataFrame()

    data = json.loads(files[-1].read_text(encoding="utf-8"))

    # Case 1 - top-level list of dicts (yearly summaries)
    if isinstance(data, list):
        return pd.DataFrame(data)

    # Case 2 - dict with nested list under a known key
    if isinstance(data, dict):
        for key in ("years", "yearly", "trades", "results", "summary",
                    "by_year", "year_stats"):
            v = data.get(key)
            if isinstance(v, list) and v and isinstance(v[0], dict):
                return pd.DataFrame(v)
        # Case 3 - flat dict
        return pd.DataFrame([data])

    return pd.DataFrame()


# ============================================================
# renderers
# ============================================================

def _clean_cell(v):
    """Convert any non-scalar to a JSON string so openpyxl can write it."""
    if v is None:
        return None
    if isinstance(v, (str, int, float, bool, date, datetime)):
        if isinstance(v, float) and pd.isna(v):
            return None
        return v
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        f = float(v)
        return None if pd.isna(f) else f
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(v, default=str)
    return str(v)


def freeze_header(ws):
    ws.freeze_panes = "A2"
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.fill = PatternFill("solid", fgColor=HDR_BG)
        cell.font = Font(color=HDR_FG, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def autosize(ws, min_w=10, max_w=30):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, longest + 2))


def band(ws, first, last, ncols):
    for r in range(first, last + 1):
        if (r - first) % 2 == 1:
            for c in range(1, ncols + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=BAND_BG)


def write_table(ws, df: pd.DataFrame,
                number_fmt: dict | None = None,
                scale_cols: list | None = None,
                color_col: str | None = None):
    if df.empty:
        ws.append(["no data"])
        freeze_header(ws)
        return

    ws.append([str(c) for c in df.columns])
    for _, row in df.iterrows():
        ws.append([_clean_cell(v) for v in row.values])

    freeze_header(ws)
    n = ws.max_row

    if number_fmt:
        for col_name, fmt in number_fmt.items():
            if col_name in df.columns:
                idx = list(df.columns).index(col_name) + 1
                for r in range(2, n + 1):
                    ws.cell(r, idx).number_format = fmt

    band(ws, 2, n, len(df.columns))

    if scale_cols:
        for col_name in scale_cols:
            if col_name in df.columns:
                idx = list(df.columns).index(col_name) + 1
                letter = get_column_letter(idx)
                ws.conditional_formatting.add(
                    f"{letter}2:{letter}{n}",
                    ColorScaleRule(start_type="min", start_color="F8696B",
                                   mid_type="num", mid_value=0,
                                   mid_color="FFFFFF",
                                   end_type="max", end_color="63BE7B"))

    if color_col and color_col in df.columns:
        idx = list(df.columns).index(color_col) + 1
        letter = get_column_letter(idx)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{n}",
            CellIsRule(operator="greaterThan", formula=["0"],
                       fill=PatternFill("solid", fgColor=UP_BG),
                       font=Font(color=UP_TEXT, bold=True)))
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{n}",
            CellIsRule(operator="lessThan", formula=["0"],
                       fill=PatternFill("solid", fgColor=DOWN_BG),
                       font=Font(color=DOWN_TEXT, bold=True)))

    autosize(ws)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{n}"
    ws.sheet_view.showGridLines = False


def add_line_chart(ws, title, df, cols, anchor=None, width=22, height=8):
    if df.empty or not cols:
        return
    n = ws.max_row
    col_idx = [list(df.columns).index(c) + 1 for c in cols if c in df.columns]
    if not col_idx:
        return
    chart = LineChart()
    chart.title = title
    chart.height = height
    chart.width  = width
    data = Reference(ws, min_col=min(col_idx), max_col=max(col_idx),
                     min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor or f"A{n+3}")


# ============================================================
# sheets
# ============================================================

def sheet_ohlcv(wb):
    df = src_ohlcv()
    ws = wb.create_sheet("OHLCV")
    if df.empty:
        write_table(ws, df); return
    out = df.reset_index()[["date","open","high","low","close","volume",
                            "return","range","dir"]].copy()
    out["date"] = pd.to_datetime(out["date"]).dt.date
    write_table(ws, out,
        number_fmt={"open":FMT_PRICE,"high":FMT_PRICE,"low":FMT_PRICE,
                    "close":FMT_PRICE,"volume":FMT_INT,
                    "return":FMT_PCT,"range":FMT_PRICE},
        scale_cols=["return"])
    add_line_chart(ws, f"{SYMBOL} - close", out, ["close"])


def sheet_indicators(wb):
    df = src_indicators()
    ws = wb.create_sheet("Indicators")
    if df.empty:
        write_table(ws, df); return
    out = df.reset_index()[["date","close","sma_20","sma_50","sma_200",
                            "ema_12","ema_26","macd","macd_sig",
                            "macd_hist","rsi_14","atr_14"]].copy()
    out["date"] = pd.to_datetime(out["date"]).dt.date
    write_table(ws, out,
        number_fmt={c: FMT_PRICE for c in
                    ["close","sma_20","sma_50","sma_200","ema_12","ema_26",
                     "macd","macd_sig","macd_hist","rsi_14","atr_14"]},
        scale_cols=["macd_hist"])
    n = ws.max_row
    ws.conditional_formatting.add(
        f"J2:J{n}",
        CellIsRule(operator="greaterThan", formula=["70"],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(
        f"J2:J{n}",
        CellIsRule(operator="lessThan", formula=["30"],
                   fill=PatternFill("solid", fgColor="C6EFCE")))
    add_line_chart(ws, "Close / SMA20 / SMA50 / SMA200",
                   out, ["close","sma_20","sma_50","sma_200"])


def sheet_stats(wb):
    df = src_stats()
    ws = wb.create_sheet("Stats")
    if df.empty:
        write_table(ws, df); return
    write_table(ws, df)
    for r in range(2, ws.max_row + 1):
        k = ws.cell(r,1).value
        v = ws.cell(r,2).value
        c = ws.cell(r,2)
        if k in ("Start price","End price","Max close","Min close"):
            c.number_format = FMT_PRICE
        elif k in ("Total return","Annualized vol","Max drawdown",
                   "Best day","Worst day"):
            c.number_format = FMT_PCT
        elif k == "Sharpe (0 rf)":
            c.number_format = "0.00"
        elif k == "Avg daily volume":
            c.number_format = FMT_INT
        if k in ("Total return","Best day") and isinstance(v,(int,float)):
            c.fill = PatternFill("solid", fgColor=UP_BG)
        elif k in ("Max drawdown","Worst day") and isinstance(v,(int,float)):
            c.fill = PatternFill("solid", fgColor=DOWN_BG)


def sheet_monthly(wb):
    df = src_monthly()
    ws = wb.create_sheet("Monthly")
    if df.empty:
        write_table(ws, df); return
    write_table(ws, df,
        number_fmt={c: "+0.0%;-0.0%;0.0%" for c in df.columns if c != "year"})
    n = ws.max_row
    last_col = get_column_letter(ws.max_column)
    ws.conditional_formatting.add(
        f"B2:{last_col}{n}",
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=UP_BG),
                   font=Font(color=UP_TEXT, bold=True)))
    ws.conditional_formatting.add(
        f"B2:{last_col}{n}",
        CellIsRule(operator="lessThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=DOWN_BG),
                   font=Font(color=DOWN_TEXT, bold=True)))


def sheet_funding(wb):
    df = src_funding()
    ws = wb.create_sheet("Funding")
    write_table(ws, df,
        number_fmt={"rate":"+0.0000%;-0.0000%"},
        scale_cols=["rate"])


def sheet_open_interest(wb):
    df = src_open_interest()
    ws = wb.create_sheet("OpenInterest")
    write_table(ws, df, number_fmt={"oi":FMT_INT,"oi_value":FMT_INT})


def sheet_fear_greed(wb):
    df = src_fear_greed()
    ws = wb.create_sheet("FearGreed")
    write_table(ws, df, scale_cols=["value"])


def sheet_options(wb):
    df = src_option_chain()
    ws = wb.create_sheet("OptionChain")
    write_table(ws, df,
        number_fmt={"strike":"#,##0","mark_price":"#,##0.0000",
                    "bid":"#,##0.0000","ask":"#,##0.0000",
                    "iv":"0.00%","oi":FMT_INT,"volume":FMT_INT})


def sheet_forecasts(wb):
    df = src_forecast_dashboard()
    ws = wb.create_sheet("Forecasts")
    write_table(ws, df,
        number_fmt={"price":FMT_PRICE,"target":FMT_PRICE,
                    "pred_ret":FMT_PCT,"confidence":"0.00"})
    if not df.empty and "action" in df.columns:
        n = ws.max_row
        idx = list(df.columns).index("action") + 1
        letter = get_column_letter(idx)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{n}",
            CellIsRule(operator="equal", formula=['"UP"'],
                       fill=PatternFill("solid", fgColor=UP_BG),
                       font=Font(color=UP_TEXT, bold=True)))
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{n}",
            CellIsRule(operator="equal", formula=['"DOWN"'],
                       fill=PatternFill("solid", fgColor=DOWN_BG),
                       font=Font(color=DOWN_TEXT, bold=True)))


def sheet_ml(wb):
    df = src_ml_leaderboard()
    ws = wb.create_sheet("ML_Leaderboard")
    write_table(ws, df,
        number_fmt={"accuracy":"0.00%","precision":"0.00%",
                    "recall":"0.00%","f1":"0.00%","n_test":FMT_INT},
        scale_cols=["accuracy"])


def sheet_backtest(wb):
    df = src_backtest()
    ws = wb.create_sheet("Backtest")
    if df.empty:
        write_table(ws, df); return

    # ensure every cell is scalar
    df = df.copy()
    for col in df.columns:
        df[col] = df[col].map(_clean_cell)

    number_fmt = {}
    if "hit" in df.columns:          number_fmt["hit"] = "0.00%"
    if "expectancy" in df.columns:   number_fmt["expectancy"] = "+0.00%;-0.00%"
    if "total_return" in df.columns: number_fmt["total_return"] = "+0.00%;-0.00%"
    if "n" in df.columns:            number_fmt["n"] = FMT_INT
    if "year" in df.columns:         number_fmt["year"] = "0"

    scale = ["total_return"] if "total_return" in df.columns else None

    write_table(ws, df, number_fmt=number_fmt, scale_cols=scale)


def sheet_inventory(wb):
    df = src_inventory()
    ws = wb.create_sheet("Inventory")
    write_table(ws, df, number_fmt={"rows":FMT_INT})


# ============================================================
# registry + profiles
# ============================================================

SHEETS = {
    "ohlcv":         sheet_ohlcv,
    "indicators":    sheet_indicators,
    "stats":         sheet_stats,
    "monthly":       sheet_monthly,
    "funding":       sheet_funding,
    "open_interest": sheet_open_interest,
    "fear_greed":    sheet_fear_greed,
    "options":       sheet_options,
    "forecasts":     sheet_forecasts,
    "ml":            sheet_ml,
    "backtest":      sheet_backtest,
    "inventory":     sheet_inventory,
}

PROFILES = {
    "full":      list(SHEETS.keys()),
    "ohlcv":     ["ohlcv", "indicators", "stats", "monthly"],
    "options":   ["options", "fear_greed"],
    "forecasts": ["forecasts"],
    "ml":        ["ml", "backtest"],
    "quality":   ["inventory"],
    "quick":     ["ohlcv", "stats"],
}


# ============================================================
# writer
# ============================================================

def build_workbook(sheet_names: list) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheet_names:
        builder = SHEETS.get(name)
        if not builder:
            print(f"  skip unknown sheet: {name}")
            continue
        t0 = time.time()
        try:
            builder(wb)
            print(f"  [OK]   {name:<15} {time.time()-t0:5.2f}s")
        except Exception as e:
            print(f"  [FAIL] {name:<15} {type(e).__name__}: {e}")
    return wb


def save_one(wb: Workbook, out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        stamp = datetime.now().strftime("%H%M%S")
        out = out.with_stem(out.stem + f"_locked_{stamp}")
        wb.save(out)
        print(f"  (file was locked - wrote {out.name} instead)")
    print(f"  wrote {out}")


def save_split(wb: Workbook, names: list, out_dir: Path) -> None:
    from openpyxl import Workbook as WB
    out_dir.mkdir(parents=True, exist_ok=True)
    for name in names:
        if name not in wb.sheetnames:
            continue
        sub = WB()
        sub.remove(sub.active)
        src = wb[name]
        dst = sub.create_sheet(name)
        for row in src.iter_rows():
            for cell in row:
                nc = dst.cell(row=cell.row, column=cell.column, value=cell.value)
                nc.number_format = cell.number_format
                if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
                    nc.fill = PatternFill("solid", fgColor=cell.fill.fgColor.rgb)
                if cell.font:
                    nc.font = Font(color=cell.font.color, bold=cell.font.bold)
        dst.freeze_panes = src.freeze_panes
        for letter, dim in src.column_dimensions.items():
            if dim.width:
                dst.column_dimensions[letter].width = dim.width
        f = out_dir / f"{name}_{TODAY}.xlsx"
        sub.save(f)
        print(f"  wrote {f}")


# ============================================================
# main
# ============================================================

def parse_args():
    p = argparse.ArgumentParser(
        description="Build Excel workbook(s) from all data sources.")
    p.add_argument("--profile", default="full", choices=list(PROFILES.keys()))
    p.add_argument("--sheet", action="append", default=None,
                   help="add a specific sheet (repeatable)")
    p.add_argument("--split", action="store_true",
                   help="write one file per sheet")
    p.add_argument("--out", default=None,
                   help="output filename (default: exports/eth_bot_<date>.xlsx)")
    p.add_argument("--list", action="store_true",
                   help="list available sheets and profiles")
    return p.parse_args()


def main() -> int:
    args = parse_args()

    if args.list:
        print("Sheets:")
        for s in SHEETS:
            print(f"  {s}")
        print("\nProfiles:")
        for p, sheets in PROFILES.items():
            print(f"  {p:<10} {', '.join(sheets)}")
        return 0

    sheets = list(PROFILES[args.profile])
    if args.sheet:
        for s in args.sheet:
            if s not in sheets:
                sheets.append(s)

    print("=" * 60)
    print(f"  EXCEL UPDATE  profile={args.profile}  sheets={len(sheets)}")
    print("=" * 60)

    wb = build_workbook(sheets)

    if args.split:
        save_split(wb, sheets, OUT_DIR)
    else:
        out = Path(args.out) if args.out else OUT_DIR / f"eth_bot_{TODAY}.xlsx"
        save_one(wb, out)

    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
