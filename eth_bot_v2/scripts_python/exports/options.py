"""Export all options data to CSV + Excel with dates.

Usage:
    python -m scripts_python.exports.options
    python -m scripts_python.exports.options --currency ETH
    python -m scripts_python.exports.options --currency ETH,BTC
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT    = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "database" / "options.db"
OUT_DIR = ROOT / "exports"


TABLES = ["option_summary", "option_chain", "implied_volatility", "greeks"]


def load(conn, table, currency=None):
    q = f"SELECT * FROM {table}"
    params = ()
    if currency:
        q += " WHERE underlying = ?"
        params = (currency,)
    q += " ORDER BY ts, expiry, strike, kind"
    return pd.read_sql_query(q, conn, params=params)


def export_csv(conn, currency, out_dir):
    written = []
    for table in TABLES:
        try:
            df = load(conn, table, currency)
        except Exception as e:
            print(f"    skip {table}: {e}")
            continue
        if df.empty:
            print(f"    skip {table}: empty")
            continue
        path = out_dir / f"options_{currency}_{table}.csv"
        df.to_csv(path, index=False)
        print(f"    csv   {path.name:<45} {len(df):>6} rows")
        written.append(path)
    return written


def export_excel(conn, currency, out_dir):
    path = out_dir / f"options_{currency}.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    any_written = False
    for table in TABLES:
        try:
            df = load(conn, table, currency)
        except Exception:
            continue
        if df.empty:
            continue

        ws = wb.create_sheet(table[:31])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append(list(row))

        # header style
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = ws.cell(row=2, column=1)

        # autosize
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = max(8, min(22, longest + 2))

        any_written = True

    if any_written:
        wb.save(path)
        print(f"    xlsx  {path.name}")
        return path
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--currency", type=str, default="ETH",
                    help="ETH, BTC, or comma-separated")
    args = ap.parse_args()

    currencies = [c.strip().upper() for c in args.currency.split(",") if c.strip()]

    if not DB_PATH.exists():
        print(f"DB not found: {DB_PATH}")
        print("Run: python database/init_options_db.py")
        return 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 60)
    print("  EXPORT OPTIONS")
    print("=" * 60)
    print(f"  currencies : {', '.join(currencies)}")
    print(f"  db         : {DB_PATH}")
    print(f"  out dir    : {OUT_DIR}")
    print()

    conn = sqlite3.connect(DB_PATH)
    try:
        all_written = []
        for cur in currencies:
            print(f"[{cur}]")
            all_written += export_csv(conn, cur, OUT_DIR)
            xlsx = export_excel(conn, cur, OUT_DIR)
            if xlsx:
                all_written.append(xlsx)
            print()

        if all_written:
            print("=" * 60)
            print("  FILES WRITTEN")
            print("=" * 60)
            for p in all_written:
                size = p.stat().st_size
                print(f"  {p.name:<48} {size:>10,} bytes")
        else:
            print("no files written (all tables empty?)")
        return 0
    finally:
        conn.close()


if __name__ == "__main__":
    sys.exit(main())