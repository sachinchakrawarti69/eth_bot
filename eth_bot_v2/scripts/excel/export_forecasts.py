"""Export all forecast JSONs to one Excel file."""
from __future__ import annotations
import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from scripts.excel.common import style_header, autosize, safe_close_excel, UP_FILL, DOWN_FILL

ROOT         = Path(__file__).resolve().parents[1]
FORECAST_DIR = ROOT / "reports" / "forecasts"
OUT_DIR      = ROOT / "exports"
STAMP        = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT          = OUT_DIR / f"ETHUSDT_forecast_{STAMP}.xlsx"


def _row(file_name, d, generated_at=""):
    return {
        "file":         file_name,
        "generated_at": generated_at or d.get("as_of", d.get("generated_at","")),
        "symbol":       d.get("symbol","ETHUSDT"),
        "action":       d.get("action", d.get("direction","")),
        "reason":       d.get("reason",""),
        "price":        d.get("price", d.get("last_close","")),
        "pred_ret":     d.get("pred_ret",""),
        "pred_close":   d.get("pred_close",""),
        "confidence":   d.get("confidence", d.get("dir_acc", d.get("dir_accuracy",""))),
    }


def flatten(file_name, data):
    if isinstance(data, list):
        return [_row(file_name, i) for i in data if isinstance(i, dict)]
    if isinstance(data, dict):
        sig = data.get("signal", data)
        return [_row(file_name, sig, data.get("generated_at",""))]
    return []


def build():
    if not FORECAST_DIR.exists():
        print(f"  no forecast dir: {FORECAST_DIR}")
        return None
    files = sorted(FORECAST_DIR.glob("*.json"))
    if not files:
        print("  no forecast json files")
        return None

    rows = []
    for fp in files:
        try:
            data = json.loads(fp.read_text(encoding="utf-8"))
        except Exception:
            continue
        rows.extend(flatten(fp.name, data))

    if not rows:
        print("  no rows after flatten")
        return None

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    safe_close_excel()
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Forecasts")
    ws.append(["file","generated_at","symbol","action","reason",
               "price","pred_ret","pred_close","confidence"])
    for r in rows:
        ws.append([r["file"], r["generated_at"], r["symbol"], r["action"],
                   r["reason"], r["price"], r["pred_ret"],
                   r["pred_close"], r["confidence"]])
    style_header(ws)
    n = ws.max_row
    for i in range(2, n + 1):
        a = ws.cell(i, 4).value
        if a in ("LONG","UP"):
            ws.cell(i,4).fill = PatternFill("solid", fgColor=UP_FILL)
            ws.cell(i,4).font = Font(color="006100", bold=True)
        elif a in ("SHORT","DOWN"):
            ws.cell(i,4).fill = PatternFill("solid", fgColor=DOWN_FILL)
            ws.cell(i,4).font = Font(color="9C0006", bold=True)
        else:
            ws.cell(i,4).fill = PatternFill("solid", fgColor="E7E6E6")
    autosize(ws, max_w=40)
    ws.freeze_panes = ws.cell(row=2, column=1)
    wb.save(OUT)
    print(f"  wrote {OUT}  ({len(rows)} rows from {len(files)} files)")
    return OUT


if __name__ == "__main__":
    build()
