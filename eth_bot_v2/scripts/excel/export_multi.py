"""Multi-symbol Excel export: prices, stats, correlations.

    python -m excel.export_multi
"""
from __future__ import annotations

import sqlite3
from datetime import datetime, timezone, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scripts.excel.common import (style_header, autosize, safe_close_excel,
                          UP_FILL, DOWN_FILL, UP_TEXT, DOWN_TEXT)


ROOT    = Path(__file__).resolve().parents[1]
DB      = ROOT / "database" / "market.db"
OUT_DIR = ROOT / "exports"

SYMBOLS    = ["BTCUSDT", "ETHUSDT", "SOLUSDT", "BNBUSDT"]
TIMEFRAME  = "1d"
LOOKBACK   = 365   # days back from "now"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT   = OUT_DIR / f"multi_symbol_{STAMP}.xlsx"


# ------------------------------------------------------------------
# data
# ------------------------------------------------------------------

def load_all():
    """Return dict {symbol: DataFrame} for the last LOOKBACK days."""
    now_ms = int(datetime.now(timezone.utc).timestamp() * 1000)
    cutoff_ms = now_ms - LOOKBACK * 86_400_000

    out = {}
    with sqlite3.connect(DB) as c:
        for sym in SYMBOLS:
            df = pd.read_sql_query(
                "SELECT open_time, open, high, low, close, volume "
                "FROM ohlcv WHERE symbol=? AND timeframe=? AND open_time >= ? "
                "ORDER BY open_time",
                c, params=(sym, TIMEFRAME, cutoff_ms))
            if df.empty:
                continue
            df["date"] = pd.to_datetime(df["open_time"], unit="ms", utc=True).dt.tz_convert(None)
            df = df.drop(columns=["open_time"]).set_index("date")
            df["return"] = df["close"].pct_change()
            out[sym] = df
    return out


def build_price_frame(data):
    """Wide DataFrame of closes, one column per symbol."""
    frames = {sym: df["close"] for sym, df in data.items()}
    return pd.concat(frames, axis=1).dropna()


def build_return_frame(data):
    """Wide DataFrame of daily returns."""
    frames = {sym: df["close"].pct_change() for sym, df in data.items()}
    return pd.concat(frames, axis=1).dropna()


def stats_for(returns):
    """Compute summary stats for a single return series."""
    r = returns.dropna()
    cum = (1 + r).cumprod()
    dd = (cum / cum.cummax() - 1).min() if len(cum) else 0.0
    ann_vol = float(r.std(ddof=0) * np.sqrt(365))
    ann_ret = float((1 + r.mean()) ** 365 - 1) if len(r) else 0.0
    sharpe = float(r.mean() / r.std(ddof=0) * np.sqrt(365)) if r.std(ddof=0) else 0.0
    return {
        "days":         int(len(r)),
        "total_return": float(cum.iloc[-1] - 1) if len(cum) else 0.0,
        "ann_return":   ann_ret,
        "ann_vol":      ann_vol,
        "sharpe":       sharpe,
        "max_dd":       float(dd),
        "best_day":     float(r.max()) if len(r) else 0.0,
        "worst_day":    float(r.min()) if len(r) else 0.0,
        "pos_days":     float((r > 0).mean()) if len(r) else 0.0,
    }


# ------------------------------------------------------------------
# sheets
# ------------------------------------------------------------------

def sheet_prices(wb, prices):
    """Close prices for all symbols, one column each."""
    ws = wb.create_sheet("Prices")
    ws.append(["Date"] + list(prices.columns))
    for d, row in prices.iterrows():
        ws.append([d.date()] + [float(v) if pd.notna(v) else None for v in row])
    style_header(ws)
    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).number_format = "#,##0.00"
    autosize(ws)
    ws.freeze_panes = ws.cell(row=2, column=1)

    # chart
    chart = LineChart()
    chart.title = f"Close price - last {LOOKBACK} days"
    chart.height, chart.width = 12, 28
    data = Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{n+3}")


def sheet_returns(wb, returns):
    """Daily returns per symbol with color scale."""
    ws = wb.create_sheet("Returns")
    ws.append(["Date"] + list(returns.columns))
    for d, row in returns.iterrows():
        ws.append([d.date()] + [float(v) if pd.notna(v) else None for v in row])
    style_header(ws)
    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).number_format = "+0.00%;-0.00%"

    last_col = get_column_letter(ws.max_column)
    ws.conditional_formatting.add(
        f"B2:{last_col}{n}",
        ColorScaleRule(start_type="num", start_value=-0.05, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=0.05, end_color="63BE7B"))
    autosize(ws)
    ws.freeze_panes = ws.cell(row=2, column=1)


def sheet_normalized(wb, prices):
    """Every symbol rebased to 100 at start - relative performance."""
    ws = wb.create_sheet("Normalized")
    norm = prices / prices.iloc[0] * 100
    ws.append(["Date"] + list(norm.columns))
    for d, row in norm.iterrows():
        ws.append([d.date()] + [float(v) if pd.notna(v) else None for v in row])
    style_header(ws)
    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).number_format = "#,##0.0"

    chart = LineChart()
    chart.title = "Normalized to 100 at start"
    chart.height, chart.width = 12, 28
    data = Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{n+3}")
    autosize(ws)


def sheet_stats(wb, data):
    """Stats table - one row per symbol."""
    ws = wb.create_sheet("Stats")
    cols = ["symbol", "days", "total_return", "ann_return", "ann_vol",
            "sharpe", "max_dd", "best_day", "worst_day", "pos_days"]
    ws.append(cols)

    rows = []
    for sym, df in data.items():
        s = stats_for(df["return"])
        rows.append([sym, s["days"], s["total_return"], s["ann_return"],
                     s["ann_vol"], s["sharpe"], s["max_dd"],
                     s["best_day"], s["worst_day"], s["pos_days"]])

    rows.sort(key=lambda r: r[3], reverse=True)   # by ann_return
    for r in rows:
        ws.append(r)

    style_header(ws)
    n = ws.max_row

    # number formats
    for r in range(2, n + 1):
        for c in (3, 4, 5, 7, 8, 9, 10):
            ws.cell(r, c).number_format = "+0.00%;-0.00%"
        ws.cell(r, 6).number_format = "0.00"

    # color scales
    # total return: green/red
    ws.conditional_formatting.add(
        f"C2:C{n}",
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=UP_FILL),
                   font=Font(color=UP_TEXT, bold=True)))
    ws.conditional_formatting.add(
        f"C2:C{n}",
        CellIsRule(operator="lessThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=DOWN_FILL),
                   font=Font(color=DOWN_TEXT, bold=True)))

    # sharpe scale
    ws.conditional_formatting.add(
        f"F2:F{n}",
        ColorScaleRule(start_type="min", start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="max", end_color="63BE7B"))

    # drawdown scale (all negative - bigger = redder)
    ws.conditional_formatting.add(
        f"G2:G{n}",
        ColorScaleRule(start_type="min", start_color="F8696B",
                       mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                       end_type="max", end_color="63BE7B"))

    autosize(ws, max_w=18)
    ws.sheet_view.showGridLines = False


def sheet_correlations(wb, returns):
    """Correlation matrix + rank."""
    ws = wb.create_sheet("Correlation")

    corr = returns.corr()
    syms = list(corr.columns)

    # matrix
    ws.append(["corr"] + syms)
    for i, sym in enumerate(syms, start=1):
        ws.append([sym] + [float(corr.loc[sym, s]) for s in syms])

    style_header(ws)
    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(2, n + 1):
            ws.cell(r, c).number_format = "0.00"

    # heatmap on the correlation block
    last_col = get_column_letter(n)
    ws.conditional_formatting.add(
        f"B2:{last_col}{n}",
        ColorScaleRule(start_type="num", start_value=0.0, start_color="FFFFFF",
                       mid_type="num", mid_value=0.5, mid_color="FFEB9C",
                       end_type="num", end_value=1.0, end_color="C6EFCE"))
    autosize(ws)
    ws.freeze_panes = ws.cell(row=2, column=2)

    # ranked pairs
    ws.append([])
    ws.append(["Correlation with BTC"])
    btc_col = "BTCUSDT" if "BTCUSDT" in corr.columns else syms[0]
    pairs = [(s, float(corr.loc[btc_col, s])) for s in syms if s != btc_col]
    pairs.sort(key=lambda x: -x[1])
    ws.append(["symbol", f"corr_vs_{btc_col}"])
    for s, c in pairs:
        ws.append([s, c])
    autosize(ws)
    ws.sheet_view.showGridLines = False


def sheet_vs_btc(wb, prices):
    """Performance ratio of each alt vs BTC."""
    ws = wb.create_sheet("vs BTC")
    btc_col = "BTCUSDT" if "BTCUSDT" in prices.columns else prices.columns[0]
    ws.append(["Date"] + [f"{s}/{btc_col}" for s in prices.columns if s != btc_col])

    btc = prices[btc_col]
    for d, row in prices.iterrows():
        vals = [float(row[s] / btc.loc[d]) if pd.notna(row[s]) and btc.loc[d] else None
                for s in prices.columns if s != btc_col]
        ws.append([d.date()] + vals)

    style_header(ws)
    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).number_format = "0.00000"

    chart = LineChart()
    chart.title = f"Alt/BTC ratios"
    chart.height, chart.width = 12, 28
    data = Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{n+3}")
    autosize(ws)


# ------------------------------------------------------------------
# main
# ------------------------------------------------------------------

def build():
    if not DB.exists():
        print(f"  DB not found: {DB}")
        return None

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    safe_close_excel()

    print(f"  cutoff: last {LOOKBACK} days from now")
    print(f"  symbols: {', '.join(SYMBOLS)}")

    data = load_all()
    if not data:
        print("  no data returned")
        return None
    for sym, df in data.items():
        print(f"    {sym:<10} {len(df)} rows  {df.index.min().date()} -> {df.index.max().date()}")

    prices  = build_price_frame(data)
    returns = build_return_frame(data)
    print(f"  aligned: {len(prices)} rows")

    wb = Workbook()
    wb.remove(wb.active)

    sheet_stats(wb, data)
    sheet_correlations(wb, returns)
    sheet_normalized(wb, prices)
    sheet_vs_btc(wb, prices)
    sheet_prices(wb, prices)
    sheet_returns(wb, returns)

    wb.save(OUT)
    print(f"  wrote {OUT}")
    return OUT


if __name__ == "__main__":
    build()
